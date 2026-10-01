"""Live dump of a design-doc workbook with openpyxl (the validated fallback / standalone route).

Writes, into OUT_DIR, one <prefix>_<sheet>.txt per visible sheet in the shared-dump format
([row,col]=value tokens, " | " separated, one line per row, empty cells skipped), with struck and
gray content removed and partially-struck cells reduced to their live text, plus
_DELETED_DIGEST.txt (one line per removed/partial cell: "<sheet> [r,c] DEL|GRAY|PART: text").
Skips 詳細設計* / *画面ｲﾒｰｼﾞ* sheets (records their size), and hidden sheets unless --hidden.

usage: python live_dump.py <workbook.xlsx> <out_dir> [--prefix P] [--sheets REGEX] [--hidden]
"""
import argparse, datetime, os, re, sys, warnings
import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
warnings.filterwarnings("ignore")

def is_gray(color):
    try:
        if color is None or color.type != "rgb" or not color.rgb: return False
        h = str(color.rgb)[-6:]
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
        return r == g == b and 80 < r < 220
    except Exception:
        return False

def fmt(v):
    if isinstance(v, datetime.datetime):
        d = v - datetime.datetime(1899, 12, 30); n = d.days + d.seconds / 86400
        return str(int(n)) if n == int(n) else repr(n)
    if isinstance(v, float) and v == int(v): return str(int(v))
    return str(v)

def classify(c):
    """-> (live_text or None, kind or None, removed_text or None)"""
    v = c.value
    if v is None: return None, None, None
    cf = c.font
    struck = bool(cf and cf.strike); gray = bool(cf and is_gray(cf.color))
    if isinstance(v, CellRichText):
        live = ""; raw = ""
        for part in v:
            if isinstance(part, TextBlock):          # its own rPr fully overrides the cell font
                f = part.font; t = part.text
                dead = bool(f and (f.strike or is_gray(f.color)))
            else:                                    # bare str inherits the cell font
                t = part; dead = struck or gray
            raw += t
            if not dead: live += t
        if not live.strip():
            return None, "DEL", raw
        if live != raw:
            return live, "PART", f"raw='{raw}' live='{live}'"
        return live, None, None
    s = fmt(v)
    if not s.strip(): return None, None, None
    if struck: return None, "DEL", s
    if gray: return None, "GRAY", s
    return s, None, None

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook"); ap.add_argument("out_dir")
    ap.add_argument("--prefix"); ap.add_argument("--sheets"); ap.add_argument("--hidden", action="store_true")
    a = ap.parse_args()
    os.makedirs(a.out_dir, exist_ok=True)
    prefix = a.prefix or os.path.basename(a.workbook).split("_")[0]
    wb = openpyxl.load_workbook(a.workbook, data_only=True, rich_text=True)
    digest = []; summary = []
    for ws in wb.worksheets:
        n = ws.title
        if ws.sheet_state != "visible" and not a.hidden: continue
        if a.sheets and not re.search(a.sheets, n): continue
        if n.startswith("詳細設計") or "画面ｲﾒｰｼﾞ" in n:
            summary.append(f"{n}\t{ws.max_row}x{ws.max_column}\tskipped"); continue
        lines = []; nd = npart = 0
        for row in ws.iter_rows():
            parts = []
            for c in row:
                live, kind, removed = classify(c)
                if live is not None: parts.append(f"[{c.row},{c.column}]={live}")
                if kind:
                    digest.append(f"{n} [{c.row},{c.column}] {kind}: {removed}")
                    if kind == "PART": npart += 1
                    else: nd += 1
            if parts: lines.append(" | ".join(parts))
        safe = "".join("_" if ch in '\\/:*?"<>|' else ch for ch in n)
        open(os.path.join(a.out_dir, f"{prefix}_{safe}.txt"), "w", encoding="utf-8").write("\n".join(lines))
        summary.append(f"{n}\t{ws.max_row}x{ws.max_column}\tdead={nd}\tpartial={npart}")
    header = ("# _DELETED_DIGEST (openpyxl live dump): one line per removed or partially-struck cell.\n"
              "# An empty list below means nothing on the dumped sheets was struck or gray.\n")
    open(os.path.join(a.out_dir, "_DELETED_DIGEST.txt"), "w", encoding="utf-8").write(header + "\n".join(digest))
    print("\n".join(summary))

if __name__ == "__main__":
    main()
